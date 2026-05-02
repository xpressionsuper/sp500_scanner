"""
report/report_exporter.py
==========================
스캔 결과를 xlsx(다중 시트) 및 PNG 이미지로 저장합니다.
"""

import sys
from datetime import datetime

import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config.settings import Config


class ReportExporter:

    def __init__(self, cfg: Config = None):
        self.cfg       = cfg or Config()
        self.today_str = datetime.today().strftime("%Y년 %m월 %d일")
        _thin          = Side(style="thin", color="CCCCCC")
        self._border   = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

    # ══════════════════════════════════════════════════════════════════
    # 공용 스타일 헬퍼
    # ══════════════════════════════════════════════════════════════════
    def _hfill(self, color: str) -> PatternFill:
        return PatternFill("solid", fgColor=color)

    def _center(self, wrap: bool = False) -> Alignment:
        return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

    def _write_title(self, ws, n_cols: int, title: str, subtitle: str) -> None:
        last = get_column_letter(n_cols)
        ws.merge_cells(f"A1:{last}1")
        tc = ws["A1"]
        tc.value = title
        tc.font      = Font(name="Arial", bold=True, size=14, color=self.cfg.C_GOLD)
        tc.fill      = self._hfill(self.cfg.C_TITLE_BG)
        tc.alignment = self._center()
        ws.row_dimensions[1].height = 30

        ws.merge_cells(f"A2:{last}2")
        sc = ws["A2"]
        sc.value     = subtitle
        sc.font      = Font(name="Arial", size=10, color="AAAAAA")
        sc.fill      = self._hfill(self.cfg.C_TITLE_BG)
        sc.alignment = self._center()
        ws.row_dimensions[2].height = 18

    def _write_header_row(self, ws, headers: list[str],
                          col_widths: dict, hdr_color: str) -> None:
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=3, column=ci, value=h)
            cell.font      = Font(name="Arial", bold=True, size=10,
                                  color=self.cfg.C_HEADER_FG)
            cell.fill      = self._hfill(hdr_color)
            cell.alignment = self._center(wrap=True)
            cell.border    = self._border
            ws.column_dimensions[get_column_letter(ci)].width = col_widths.get(h, 13)
        ws.row_dimensions[3].height = 26

    def _write_empty_sheet(self, ws, title: str, n_cols: int = 9) -> None:
        last = get_column_letter(n_cols)
        widths = [10, 10, 14, 13, 13, 13, 16, 12, 10]
        for i in range(1, n_cols + 1):
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]

        ws.merge_cells(f"A1:{last}1")
        tc = ws["A1"]
        tc.value     = title
        tc.font      = Font(name="Arial", bold=True, size=14, color=self.cfg.C_GOLD)
        tc.fill      = self._hfill(self.cfg.C_TITLE_BG)
        tc.alignment = self._center()
        ws.row_dimensions[1].height = 30

        ws.merge_cells(f"A2:{last}2")
        sc = ws["A2"]
        sc.value     = "시그널 종목 없음"
        sc.font      = Font(name="Arial", size=11, color="AAAAAA")
        sc.fill      = self._hfill(self.cfg.C_TITLE_BG)
        sc.alignment = self._center()
        ws.row_dimensions[2].height = 24

    # ══════════════════════════════════════════════════════════════════
    # 시트별 write 메서드
    # ══════════════════════════════════════════════════════════════════
    def _write_ichimoku_sheet(self, ws, df: pd.DataFrame) -> None:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A4"
        cfg = self.cfg

        if df.empty:
            self._write_empty_sheet(ws, f"📈  일목균형표 매수 시그널  |  {self.today_str}")
            return

        n_cols     = len(df.columns)
        daily_cnt  = len(df[df["타임프레임"] == "일봉"])
        weekly_cnt = len(df[df["타임프레임"] == "주봉"])
        self._write_title(
            ws, n_cols,
            f"📈  일목균형표 매수 시그널  |  {self.today_str}",
            f"총 {len(df)}개  |  일봉 {daily_cnt}개  주봉 {weekly_cnt}개  |  "
            "🟢 구름 돌파  🟡 지지 반등  |  9-26-52"
        )

        col_widths = {
            "티커":12,"타임프레임":10,"시그널 유형":14,"종가 ($)":13,
            "구름 거리 (%)":13,"52주 위치 (%)":13,"기준선 기울기 (%)":16,
            "구름 색":12,"전환>기준":10,
        }
        headers = list(df.columns)
        self._write_header_row(ws, headers, col_widths, cfg.C_HEADER_BG)

        for ri, (_, row) in enumerate(df.iterrows(), 4):
            tf     = str(row.get("타임프레임", ""))
            sig    = str(row.get("시그널 유형", ""))
            row_bg = (cfg.C_BO if "돌파" in sig else cfg.C_PB) if tf == "일봉" \
                else ("D6EAF8" if "돌파" in sig else "EAFAF1")

            for ci, col in enumerate(headers, 1):
                val  = row[col]
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.fill = self._hfill(row_bg); cell.border = self._border
                cell.alignment = self._center(); cell.font = Font(name="Arial", size=10)

                if col == "티커":
                    cell.font = Font(name="Arial", size=10, bold=True)
                elif col == "타임프레임":
                    c = cfg.C_DAILY_HDR if val == "일봉" else cfg.C_WEEKLY_HDR
                    cell.font = Font(name="Arial", size=10, bold=True, color=c)
                elif col == "시그널 유형":
                    c = "1A5276" if "돌파" in str(val) else "7D6608"
                    cell.font = Font(name="Arial", size=10, bold=True, color=c)
                elif col == "구름 거리 (%)":
                    try:
                        v = float(str(val).replace("%", ""))
                        cell.font = Font(name="Arial", size=10,
                                         color=cfg.C_GREEN if v >= 0 else cfg.C_RED)
                    except: pass
                elif col == "52주 위치 (%)":
                    try:
                        cell.font = Font(name="Arial", size=10,
                                         color=cfg.C_GREEN if float(val) >= 40 else "FF8C00")
                    except: pass
                elif col == "기준선 기울기 (%)":
                    try:
                        cell.font = Font(name="Arial", size=10,
                                         color=cfg.C_GREEN if float(val) >= 0 else cfg.C_RED)
                    except: pass
                elif col == "구름 색":
                    cell.font = Font(name="Arial", size=10,
                                     color=cfg.C_GREEN if "녹" in str(val) else cfg.C_RED)
                elif col == "전환>기준":
                    cell.font = Font(name="Arial", size=10,
                                     color=cfg.C_GREEN if val == "O" else cfg.C_RED,
                                     bold=(val == "O"))
            ws.row_dimensions[ri].height = 20

        ws.auto_filter.ref = f"A3:{get_column_letter(n_cols)}{3 + len(df)}"

    def _write_cup_sheet(self, ws, df: pd.DataFrame) -> None:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A4"
        cfg = self.cfg

        if df.empty:
            self._write_empty_sheet(ws, f"🏆  컵앤핸들 패턴  |  {self.today_str}")
            return

        n_cols     = len(df.columns)
        daily_cnt  = len(df[df["타임프레임"] == "일봉"]) if "타임프레임" in df.columns else 0
        weekly_cnt = len(df[df["타임프레임"] == "주봉"]) if "타임프레임" in df.columns else 0
        self._write_title(
            ws, n_cols,
            f"🏆  컵앤핸들 패턴  |  {self.today_str}",
            f"총 {len(df)}개  |  일봉 {daily_cnt}개  주봉 {weekly_cnt}개  |  "
            "윌리엄 오닐 기준  |  lookback 120봉"
        )

        col_widths = {
            "티커":10,"타임프레임":10,"시그널 유형":12,"종가 ($)":13,
            "돌파선까지 (%)":15,"컵 깊이 (%)":12,"52주 위치 (%)":13,
            "컵 왼쪽 고점":13,"컵 바닥":12,"돌파 기준선":13,
        }
        headers = list(df.columns)
        self._write_header_row(ws, headers, col_widths, cfg.C_CUP_HDR)

        for ri, (_, row) in enumerate(df.iterrows(), 4):
            tf     = str(row.get("타임프레임", ""))
            row_bg = "F5EEF8" if tf == "일봉" else "EAF0FB"

            for ci, col in enumerate(headers, 1):
                val  = row[col]
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.fill = self._hfill(row_bg); cell.border = self._border
                cell.alignment = self._center(); cell.font = Font(name="Arial", size=10)

                if col == "티커":
                    cell.font = Font(name="Arial", size=10, bold=True)
                elif col == "타임프레임":
                    c = cfg.C_DAILY_HDR if val == "일봉" else cfg.C_WEEKLY_HDR
                    cell.font = Font(name="Arial", size=10, bold=True, color=c)
                elif col == "시그널 유형":
                    cell.font = Font(name="Arial", size=10, bold=True, color=cfg.C_CUP_HDR)
                elif col == "돌파선까지 (%)":
                    try:
                        v = float(str(val).replace("%", ""))
                        cell.font = Font(name="Arial", size=10,
                                         color=cfg.C_GREEN if v <= 2
                                         else ("FF8C00" if v <= 5 else cfg.C_RED))
                    except: pass
                elif col == "컵 깊이 (%)":
                    try:
                        v = float(str(val).replace("%", ""))
                        cell.font = Font(name="Arial", size=10,
                                         color=cfg.C_GREEN if 15 <= v <= 35 else "FF8C00")
                    except: pass
                elif col == "52주 위치 (%)":
                    try:
                        cell.font = Font(name="Arial", size=10,
                                         color=cfg.C_GREEN if float(val) >= 40 else "FF8C00")
                    except: pass
            ws.row_dimensions[ri].height = 20

        ws.auto_filter.ref = f"A3:{get_column_letter(n_cols)}{3 + len(df)}"

    def _write_candle_sheet(self, ws, df: pd.DataFrame) -> None:
        ws.sheet_view.showGridLines = False
        cfg = self.cfg

        if df is None or df.empty:
            self._write_empty_sheet(ws, f"캔들 패턴 스캔  |  {self.today_str}")
            return

        ws.freeze_panes = "A4"
        n_cols     = len(df.columns)
        daily_cnt  = len(df[df["타임프레임"] == "일봉"]) if "타임프레임" in df.columns else 0
        weekly_cnt = len(df[df["타임프레임"] == "주봉"]) if "타임프레임" in df.columns else 0
        self._write_title(
            ws, n_cols,
            f"캔들 패턴 스캔  |  {self.today_str}",
            f"총 {len(df)}개  |  일봉 {daily_cnt}개  주봉 {weekly_cnt}개  |  "
            "샛별형 도지 / 관통형"
        )

        col_widths = {
            "티커":10,"타임프레임":10,"시그널 유형":14,"종가 ($)":12,
            "52주 위치 (%)":13,"시가":11,"고가":11,"저가":11,
        }
        headers = list(df.columns)
        self._write_header_row(ws, headers, col_widths, cfg.C_CANDLE_HDR)

        for ri, (_, row) in enumerate(df.iterrows(), 4):
            tf     = str(row.get("타임프레임", ""))
            row_bg = "FDF5E6" if tf == "일봉" else "E8F4FD"

            for ci, col in enumerate(headers, 1):
                val  = row[col]
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.fill = self._hfill(row_bg); cell.border = self._border
                cell.alignment = self._center(); cell.font = Font(name="Arial", size=10)

                if col == "티커":
                    cell.font = Font(name="Arial", size=10, bold=True)
                elif col == "타임프레임":
                    c = cfg.C_DAILY_HDR if val == "일봉" else cfg.C_WEEKLY_HDR
                    cell.font = Font(name="Arial", size=10, bold=True, color=c)
                elif col == "시그널 유형":
                    cell.font = Font(name="Arial", size=10, bold=True, color=cfg.C_CANDLE_HDR)
                elif col == "52주 위치 (%)":
                    try:
                        cell.font = Font(name="Arial", size=10,
                                         color=cfg.C_GREEN if float(val) >= 40 else "FF8C00")
                    except: pass
            ws.row_dimensions[ri].height = 20

        ws.auto_filter.ref = f"A3:{get_column_letter(n_cols)}{3 + len(df)}"

    def _write_backtest_sheet(self, ws, df: pd.DataFrame,
                               title_text: str, hdr_color: str) -> None:
        ws.sheet_view.showGridLines = False
        cfg = self.cfg

        if df is None or df.empty:
            self._write_empty_sheet(ws, f"{title_text}  |  {self.today_str}")
            return

        ws.freeze_panes = "A4"
        n_cols  = len(df.columns)
        pos_cnt = len(df[df["수익률 (%)"] >= 0]) if "수익률 (%)" in df.columns else 0
        neg_cnt = len(df) - pos_cnt
        n_dates = df["시그널 날짜"].nunique() if "시그널 날짜" in df.columns else "-"
        self._write_title(
            ws, n_cols,
            f"{title_text}  |  {self.today_str}",
            f"총 {len(df)}건  |  수익 {pos_cnt}건  손실 {neg_cnt}건  |  "
            f"최근 {n_dates}개 날짜"
        )

        col_widths = {
            "티커":9,"시그널 날짜":13,"타임프레임":9,"시그널 유형":13,
            "진입가 ($)":12,"현재가 ($)":12,"수익률 (%)":12,
            "구름 거리(%)":12,"구름 색":11,"전환>기준":9,
            "돌파선까지(%)":13,"컵 깊이(%)":11,
        }
        headers = list(df.columns)
        self._write_header_row(ws, headers, col_widths, hdr_color)

        for ri, (_, row) in enumerate(df.iterrows(), 4):
            try: ret_val = float(row.get("수익률 (%)", 0))
            except: ret_val = 0
            row_bg = "E8F5E9" if ret_val >= 0 else "FFEBEE"

            for ci, col in enumerate(headers, 1):
                val  = row[col]
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.fill = self._hfill(row_bg); cell.border = self._border
                cell.alignment = self._center(); cell.font = Font(name="Arial", size=9)

                if col == "수익률 (%)":
                    try:
                        v = float(val)
                        cell.font = Font(name="Arial", size=9, bold=True,
                                         color=cfg.C_GREEN if v >= 0 else cfg.C_RED)
                    except: pass
                elif col == "티커":
                    cell.font = Font(name="Arial", size=9, bold=True)
                elif col == "시그널 날짜":
                    cell.font = Font(name="Arial", size=9, color="555555")
                elif col == "시그널 유형":
                    c = ("1A5276" if "돌파" in str(val)
                         else "7D6608" if "반등" in str(val) else "6E2F8A")
                    cell.font = Font(name="Arial", size=9, bold=True, color=c)
                elif col == "구름 색":
                    cell.font = Font(name="Arial", size=9,
                                     color=cfg.C_GREEN if "녹" in str(val) else cfg.C_RED)
            ws.row_dimensions[ri].height = 18

        ws.auto_filter.ref = f"A3:{get_column_letter(n_cols)}{3 + len(df)}"

    def _write_summary_sheet(self, ws, df_i: pd.DataFrame,
                              df_c: pd.DataFrame, df_candle: pd.DataFrame) -> None:
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20
        cfg = self.cfg

        ws.merge_cells("A1:B1")
        t = ws["A1"]
        t.value = "📊  스캔 요약"
        t.font = Font(name="Arial", bold=True, size=13, color=cfg.C_GOLD)
        t.fill = self._hfill(cfg.C_TITLE_BG)
        t.alignment = self._center()
        ws.row_dimensions[1].height = 28

        def cnt(df, tf=None, sig=None):
            if df is None or df.empty: return 0
            d = df
            if tf and "타임프레임" in d.columns: d = d[d["타임프레임"] == tf]
            if sig and "시그널 유형" in d.columns:
                d = d[d["시그널 유형"].str.contains(sig, na=False)]
            return len(d)

        n_i = len(df_i) if df_i is not None else 0
        n_c = len(df_c) if df_c is not None else 0
        n_cd = len(df_candle) if df_candle is not None else 0

        rows = [
            ("분석 일시",           datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("스캔 대상",           "S&P 500 + S&P 400 MidCap"),
            ("─────────────────", "──────────────"),
            ("【일목균형표】",        ""),
            ("  일봉 시그널",         cnt(df_i, tf="일봉")),
            ("    ├ 구름 돌파",      cnt(df_i, tf="일봉", sig="돌파")),
            ("    └ 지지 반등",      cnt(df_i, tf="일봉", sig="반등")),
            ("  주봉 시그널",         cnt(df_i, tf="주봉")),
            ("    ├ 구름 돌파",      cnt(df_i, tf="주봉", sig="돌파")),
            ("    └ 지지 반등",      cnt(df_i, tf="주봉", sig="반등")),
            ("  소계",               n_i),
            ("─────────────────", "──────────────"),
            ("【컵앤핸들】",          ""),
            ("  일봉 패턴",           cnt(df_c, tf="일봉")),
            ("  주봉 패턴",           cnt(df_c, tf="주봉")),
            ("  소계",               n_c),
            ("─────────────────", "──────────────"),
            ("【캔들 패턴】",          ""),
            ("  일봉 패턴",           cnt(df_candle, tf="일봉")),
            ("  주봉 패턴",           cnt(df_candle, tf="주봉")),
            ("  소계",               n_cd),
            ("─────────────────", "──────────────"),
            ("총 시그널 합계",        n_i + n_c + n_cd),
        ]

        for ri, (label, value) in enumerate(rows, 2):
            ca  = ws.cell(row=ri, column=1, value=label)
            cb2 = ws.cell(row=ri, column=2, value=value)
            bg  = "F0F4FF" if ri % 2 == 0 else "FFFFFF"
            if "─" in label: bg = "E8E8E8"
            if "【" in label: bg = cfg.C_TITLE_BG
            for cell in [ca, cb2]:
                cell.fill = self._hfill(bg); cell.border = self._border
                cell.alignment = self._center()
                cell.font = Font(name="Arial", size=10,
                                 color="FFFFFF" if "【" in label else "000000")
            ca.font = Font(name="Arial", size=10, bold=True,
                           color="FFFFFF" if "【" in label else "000000")
            ws.row_dimensions[ri].height = 20

    # ══════════════════════════════════════════════════════════════════
    # xlsx 저장 메인
    # ══════════════════════════════════════════════════════════════════
    def save_xlsx(self, filename: str,
                  df_ichimoku: pd.DataFrame,
                  df_cup: pd.DataFrame,
                  df_candle: pd.DataFrame,
                  df_bt_ich: pd.DataFrame    = None,
                  df_bt_cup: pd.DataFrame    = None,
                  df_bt_candle: pd.DataFrame = None) -> None:
        """모든 시트를 포함한 xlsx 리포트를 저장합니다."""
        cfg = self.cfg
        wb  = Workbook()

        ws = wb.active; ws.title = "일목균형표"
        self._write_ichimoku_sheet(ws, df_ichimoku)

        ws = wb.create_sheet("컵앤핸들")
        self._write_cup_sheet(ws, df_cup)

        ws = wb.create_sheet("캔들패턴")
        self._write_candle_sheet(ws, df_candle if df_candle is not None else pd.DataFrame())

        ws = wb.create_sheet("백테스트_일목")
        self._write_backtest_sheet(
            ws, df_bt_ich if df_bt_ich is not None else pd.DataFrame(),
            "일목균형표 백테스트 (과거 시그널 → 현재 수익률)", cfg.C_HEADER_BG
        )
        ws = wb.create_sheet("백테스트_컵핸들")
        self._write_backtest_sheet(
            ws, df_bt_cup if df_bt_cup is not None else pd.DataFrame(),
            "컵앤핸들 백테스트 (과거 시그널 → 현재 수익률)", cfg.C_CUP_HDR
        )
        ws = wb.create_sheet("백테스트_캔들")
        self._write_backtest_sheet(
            ws, df_bt_candle if df_bt_candle is not None else pd.DataFrame(),
            "캔들 패턴 백테스트 (과거 시그널 → 현재 수익률)", cfg.C_CANDLE_HDR
        )
        ws = wb.create_sheet("분석요약")
        self._write_summary_sheet(ws, df_ichimoku, df_cup, df_candle)

        wb.save(filename)
        print(f"💾 저장 완료: {filename}")

    # ══════════════════════════════════════════════════════════════════
    # 이미지 저장
    # ══════════════════════════════════════════════════════════════════
    @staticmethod
    def _set_korean_font() -> None:
        import matplotlib.font_manager as fm
        candidates = {
            "win32":  ["Malgun Gothic", "맑은 고딕", "NanumGothic"],
            "darwin": ["AppleGothic", "Apple SD Gothic Neo", "NanumGothic"],
            "linux":  ["NanumGothic", "NanumBarunGothic", "UnDotum", "DejaVu Sans"],
        }
        installed = {f.name for f in fm.fontManager.ttflist}
        for font in candidates.get(sys.platform, candidates["linux"]):
            if font in installed:
                plt.rcParams["font.family"]       = font
                plt.rcParams["axes.unicode_minus"] = False
                return
        plt.rcParams["axes.unicode_minus"] = False

    @staticmethod
    def _pil_font(size: int):
        from PIL import ImageFont
        import matplotlib.font_manager as fm
        candidates = {
            "win32":  ["Malgun Gothic", "맑은 고딕", "NanumGothic"],
            "darwin": ["AppleGothic", "Apple SD Gothic Neo", "NanumGothic"],
            "linux":  ["NanumGothic", "NanumBarunGothic", "UnDotum"],
        }
        installed = {f.name: f.fname for f in fm.fontManager.ttflist}
        for name in candidates.get(sys.platform, candidates["linux"]):
            if name in installed:
                try:
                    return ImageFont.truetype(installed[name], size)
                except: pass
        try:
            return ImageFont.truetype(
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", size)
        except:
            return ImageFont.load_default()

    def save_image(self, df: pd.DataFrame, filename: str, title: str) -> None:
        """DataFrame 을 PNG 이미지로 저장합니다 (PIL 기반 직접 렌더링)."""
        from PIL import Image, ImageDraw

        self._set_korean_font()

        if df is None or df.empty:
            img  = Image.new("RGB", (600, 120), (22, 33, 62))
            draw = ImageDraw.Draw(img)
            fnt  = self._pil_font(14)
            draw.text((20, 45), f"{title}  —  해당 기간 시그널 없음",
                      font=fnt, fill=(255, 192, 0))
            img.save(filename, dpi=(180, 180))
            print(f"🖼️  이미지 저장 완료 (결과 없음): {filename}")
            return

        MAX_ROWS   = 40
        display_df = df.head(MAX_ROWS).reset_index(drop=True)
        n_rows     = len(display_df)
        n_cols     = len(display_df.columns)
        TITLE_H    = 54; HEADER_H = 36; ROW_H = 28; PAD_X = 14
        BORDER_CLR = (200, 200, 200); BG_PAGE = (22, 33, 62)

        font_hdr   = self._pil_font(12)
        font_data  = self._pil_font(11)
        font_title = self._pil_font(16)

        def text_w(text, font):
            dummy = Image.new("RGB", (1, 1))
            d = ImageDraw.Draw(dummy)
            bb = d.textbbox((0, 0), str(text), font=font)
            return bb[2] - bb[0]

        col_widths = []
        for ci, col in enumerate(display_df.columns):
            mw = text_w(col, font_hdr) + PAD_X * 2
            for ri in range(n_rows):
                mw = max(mw, text_w(str(display_df.iloc[ri, ci]), font_data) + PAD_X * 2)
            col_widths.append(max(mw, 70))

        total_w = sum(col_widths)
        total_h = TITLE_H + HEADER_H + ROW_H * n_rows + 10
        img  = Image.new("RGB", (total_w, total_h), BG_PAGE)
        draw = ImageDraw.Draw(img)

        suffix = f"  (상위 {n_rows}개)" if len(df) > MAX_ROWS else f"  (총 {n_rows}개)"
        tw = text_w(title + suffix, font_title)
        draw.text(((total_w - tw) // 2, 14), title + suffix,
                  font=font_title, fill=(255, 192, 0))
        ds = datetime.today().strftime("%Y-%m-%d")
        dw = text_w(ds, font_data)
        draw.text((total_w - dw - 10, 18), ds, font=font_data, fill=(170, 170, 170))

        y = TITLE_H; x = 0
        for ci, col in enumerate(display_df.columns):
            cw = col_widths[ci]
            draw.rectangle([x, y, x + cw - 1, y + HEADER_H - 1], fill=(26, 26, 46))
            if ci > 0: draw.line([(x, y), (x, y + HEADER_H)], fill=BORDER_CLR, width=1)
            tw2 = text_w(col, font_hdr)
            try:
                bb = draw.textbbox((0, 0), col, font=font_hdr); th = bb[3] - bb[1]
            except: th = 14
            draw.text((x + (cw - tw2) // 2, y + (HEADER_H - th) // 2),
                      col, font=font_hdr, fill=(255, 255, 255))
            x += cw
        draw.rectangle([0, y + HEADER_H - 2, total_w, y + HEADER_H], fill=(80, 80, 80))

        LIGHT = [(240, 244, 255), (255, 255, 255)]
        for ri in range(n_rows):
            y = TITLE_H + HEADER_H + ROW_H * ri; x = 0; bg = LIGHT[ri % 2]
            for ci in range(n_cols):
                cw = col_widths[ci]; val = str(display_df.iloc[ri, ci])
                draw.rectangle([x, y, x + cw - 1, y + ROW_H - 1], fill=bg)
                if ci > 0: draw.line([(x, y), (x, y + ROW_H)], fill=BORDER_CLR, width=1)
                tw2 = text_w(val, font_data)
                try:
                    bb = draw.textbbox((0, 0), val, font=font_data); th = bb[3] - bb[1]
                except: th = 13
                draw.text((x + (cw - tw2) // 2, y + (ROW_H - th) // 2),
                          val, font=font_data, fill=(30, 30, 30))
                x += cw
            draw.line([(0, y + ROW_H - 1), (total_w, y + ROW_H - 1)],
                      fill=BORDER_CLR, width=1)

        draw.rectangle([0, TITLE_H, total_w - 1, total_h - 10], outline=(80, 80, 80), width=1)
        img.save(filename, dpi=(180, 180))
        print(f"🖼️  이미지 저장 완료: {filename}  ({total_w}×{total_h}px)")
